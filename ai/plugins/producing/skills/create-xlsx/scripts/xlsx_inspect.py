"""Independent openpyxl verification and bounded workbook summaries."""
from datetime import date, datetime
import io
import math
import openpyxl
from openpyxl.utils.datetime import to_excel
from xlsx_support import DocumentError


def books(content):
    return (openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False),
            openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False))


def equal_number(actual, expected):
    return type(actual) in (int, float) and math.isclose(actual, expected, rel_tol=1e-14, abs_tol=1e-12)


def verify(content, request):
    formulas, cached = books(content)
    try:
        if formulas.sheetnames != [s['name'] for s in request['sheets']]:
            raise DocumentError(5, 'validation: worksheet names differ')
        for spec in request['sheets']:
            page, cache = formulas[spec['name']], cached[spec['name']]
            dimensions = dict(max_row=len(spec['rows']), max_col=len(spec['rows'][0]))
            for row, actual_row, cache_row in zip(spec['rows'], page.iter_rows(**dimensions), cache.iter_rows(**dimensions)):
                for expected, actual, cache_cell in zip(row, actual_row, cache_row):
                    kind, value = expected['type'], expected.get('value')
                    if kind == 'formula':
                        valid = actual.data_type == 'f' and actual.value == value and equal_number(cache_cell.value, expected.get('cached', 0))
                    elif kind == 'number':
                        valid = actual.data_type == 'n' and equal_number(actual.value, value)
                    elif kind == 'date':
                        numeric = to_excel(datetime.combine(date.fromisoformat(value), datetime.min.time()))
                        valid = (actual.value == datetime.combine(date.fromisoformat(value), datetime.min.time())
                                 or equal_number(actual.value, numeric))
                    elif kind == 'blank':
                        valid = actual.value is None
                    elif kind == 'string':
                        valid = actual.value == value and actual.data_type == 's'
                    else:
                        valid = actual.value is value and actual.data_type == 'b'
                    if not valid:
                        raise DocumentError(5, 'validation: cell type/value/formula/cache differs')
                    expected_format = request.get('formats', {}).get(expected.get('format'), {}).get('num_format')
                    if expected_format and actual.number_format != expected_format:
                        raise DocumentError(5, 'validation: number format differs')
    finally:
        formulas.close(); cached.close()


def inspect(content):
    formulas, cached = books(content)
    remaining, chars, truncated = 2000, 50000, False
    def value(item):
        nonlocal chars, truncated
        if isinstance(item, (datetime, date)):
            item = item.isoformat()
        if isinstance(item, str):
            result = item[:chars]
            chars -= len(result)
            truncated |= len(result) < len(item)
            return result
        return item
    result = []
    try:
        for name in formulas.sheetnames[:20]:
            sheet, cache = formulas[name], cached[name]
            cells = []
            # Zip streams avoids repeated random access in read-only worksheets.
            truncated |= (sheet.max_column or 0) > 100 or (sheet.max_row or 0) > 2000
            bounds = dict(max_col=min(sheet.max_column or 1, 100), max_row=min(sheet.max_row or 1, 2001))
            for row, cached_row in zip(sheet.iter_rows(**bounds), cache.iter_rows(**bounds)):
                if not remaining:
                    truncated = True
                    break
                for cell, cached_cell in zip(row, cached_row):
                    if not remaining:
                        truncated = True
                        break
                    remaining -= 1
                    item = dict(coordinate=getattr(cell, 'coordinate', None), type=cell.data_type, value=value(cell.value))
                    if cell.data_type == 'f':
                        item.update(cached=value(cached_cell.value), cache_source='unknown')
                    cells.append(item)
            result.append(dict(name=value(name), cells=cells))
        truncated |= len(formulas.sheetnames) > 20
        calculation = formulas.calculation
        return dict(sheets=result, truncated=truncated, recalculated=False,
                    calculation=dict(mode=calculation.calcMode, full_calc_on_load=calculation.fullCalcOnLoad) if calculation else None)
    finally:
        formulas.close(); cached.close()
