"""Bounded structural summaries for created presentations, without rendering."""
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def geometry(shape):
    return {key: getattr(shape, key) for key in ('left', 'top', 'width', 'height')}


def layouts():
    presentation = Presentation()
    return {'layouts': [{'name': layout.name, 'placeholders': [
        dict(idx=p.placeholder_format.idx, name=p.name, type=str(p.placeholder_format.type),
             geometry=geometry(p)) for p in layout.placeholders]} for layout in presentation.slide_layouts]}


def inspect(presentation):
    remaining, values_remaining = 50000, 10000
    truncated = False

    def excerpt(value):
        nonlocal remaining, truncated
        result = value[:remaining]
        remaining -= len(result)
        truncated |= len(result) < len(value)
        return result

    def take(values):
        nonlocal values_remaining, truncated
        values = list(values)
        result = values[:values_remaining]
        values_remaining -= len(result)
        truncated |= len(result) < len(values)
        return result

    slides = []
    for index, slide in enumerate(presentation.slides):
        if index >= 100:
            truncated = True
            break
        shapes = []
        for shape_index, shape in enumerate(slide.shapes):
            if shape_index >= 100:
                truncated = True
                break
            item = {'geometry': geometry(shape), 'type': str(shape.shape_type)}
            if shape.has_text_frame:
                item['text'] = excerpt(shape.text)
            if shape.has_table:
                item['cells'] = [excerpt(cell.text) for cell in take(
                    cell for row in shape.table.rows for cell in row.cells)]
                item['rows'], item['columns'] = len(shape.table.rows), len(shape.table.columns)
            if shape.has_chart:
                chart = shape.chart
                item['categories'] = [excerpt(c.label) for c in take(chart.plots[0].categories)]
                item['series'] = [{'name': excerpt(s.name), 'values': take(s.values)}
                                  for s in take(chart.series)]
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                item['image_sha1'] = shape.image.sha1
            shapes.append(item)
        slides.append({'slide': index + 1, 'shapes': shapes})
    return {'width': presentation.slide_width, 'height': presentation.slide_height,
            'slide_count': len(presentation.slides), 'slides': slides, 'truncated': truncated}
