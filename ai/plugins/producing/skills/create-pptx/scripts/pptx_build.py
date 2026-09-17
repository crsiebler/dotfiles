"""Original bounded editable slide construction using public python-pptx APIs."""
import io
import re
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt
from PIL import Image
from pptx_support import DocumentError, fields, number, project_path, text


def sequence(value, low, high):
    if not isinstance(value, list) or not low <= len(value) <= high:
        raise DocumentError(2, 'input: invalid list length')
    return value


def geometry(shape, width, height):
    x, y = number(shape.get('x'), 0, width), number(shape.get('y'), 0, height)
    w, h = number(shape.get('width'), .01, width), number(shape.get('height'), .01, height)
    if x + w > width + 1e-8 or y + h > height + 1e-8:
        raise DocumentError(2, 'input: shape extends beyond slide')
    return tuple(Inches(v) for v in (x, y, w, h))


def build(request, root):
    fields(request, ('width', 'height', 'slides'))
    width, height = number(request.get('width'), 1, 56), number(request.get('height'), 1, 56)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(width), Inches(height)
    for item in sequence(request.get('slides'), 1, 100):
        fields(item, ('layout', 'placeholders', 'shapes'))
        layouts = [layout for layout in presentation.slide_layouts if layout.name == item.get('layout')]
        if len(layouts) != 1:
            raise DocumentError(2, 'input: choose a unique inspected layout name')
        slide = presentation.slides.add_slide(layouts[0])
        seen = set()
        for entry in sequence(item.get('placeholders', []), 0, 100):
            fields(entry, ('idx', 'text'))
            idx = entry.get('idx')
            if type(idx) is not int or idx in seen:
                raise DocumentError(2, 'input: invalid or duplicate placeholder index')
            seen.add(idx)
            matches = [p for p in slide.placeholders if p.placeholder_format.idx == idx and p.has_text_frame]
            if len(matches) != 1:
                raise DocumentError(2, 'input: selected placeholder does not support text')
            matches[0].text = text(entry.get('text'))
        for shape in sequence(item.get('shapes'), 0, 100):
            add_shape(slide, shape, width, height, root)
    return presentation


def add_shape(slide, shape, width, height, root):
    if not isinstance(shape, dict):
        raise DocumentError(2, 'input: expected shape object')
    kind = shape.get('kind')
    extra = {'text': ('text', 'font', 'size', 'bold', 'color'), 'image': ('path',),
             'table': ('rows', 'size'), 'chart': ('categories', 'series')}
    if kind not in extra:
        raise DocumentError(2, 'input: unsupported shape kind')
    fields(shape, ('kind', 'x', 'y', 'width', 'height', *extra[kind]))
    x, y, w, h = geometry(shape, width, height)
    if kind == 'text':
        value = text(shape.get('text'))
        font = text(shape.get('font', 'Arial'), 200)
        size = number(shape.get('size', 20), 8, 96)
        bold, color = shape.get('bold', False), shape.get('color', '202020')
        if type(bold) is not bool or not isinstance(color, str) or not re.fullmatch('[0-9a-fA-F]{6}', color):
            raise DocumentError(2, 'input: invalid text formatting')
        frame = slide.shapes.add_textbox(x, y, w, h).text_frame
        frame.word_wrap = True
        frame.text = value
        for paragraph in frame.paragraphs:
            paragraph.font.name, paragraph.font.size = font, Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
    elif kind == 'image':
        content = project_path(root, shape.get('path')).read_bytes()
        try:
            with Image.open(io.BytesIO(content)) as image:
                if image.format not in ('PNG', 'JPEG') or image.width * image.height > 25000000:
                    raise ValueError('unsupported image')
                iw, ih = image.size
                image.verify()
        except Exception:
            raise DocumentError(2, 'input: expected valid PNG/JPEG of at most 25 million pixels') from None
        scale = min(w / iw, h / ih)
        fitted_w, fitted_h = int(iw * scale), int(ih * scale)
        slide.shapes.add_picture(io.BytesIO(content), x + (w - fitted_w) // 2,
                                 y + (h - fitted_h) // 2, fitted_w, fitted_h)
    elif kind == 'table':
        rows = sequence(shape.get('rows'), 1, 50)
        columns = len(sequence(rows[0], 1, 20))
        for row in rows:
            sequence(row, columns, columns)
            for value in row:
                text(value)
        size = number(shape.get('size', 16), 8, 48)
        table = slide.shapes.add_table(len(rows), columns, x, y, w, h).table
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
                for paragraph in table.cell(r, c).text_frame.paragraphs:
                    paragraph.font.size = Pt(size)
    else:
        categories = [text(v, 500) for v in sequence(shape.get('categories'), 1, 100)]
        data = CategoryChartData()
        data.categories = categories
        for series in sequence(shape.get('series'), 1, 10):
            fields(series, ('name', 'values'))
            values = [number(v, -1e100, 1e100) for v in sequence(series.get('values'), len(categories), len(categories))]
            data.add_series(text(series.get('name'), 500), values)
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, data)
