"""Bounded structural inspection; never render, convert, or save the input."""
import io
from docx import Document
from docx_support import inspect_package


def inspect_document(content):
    inspect_package(content)
    document = Document(io.BytesIO(content))
    remaining = 50000
    truncated = False

    def excerpt(value):
        nonlocal remaining, truncated
        length = min(2000, remaining)
        result = value[:length]
        remaining -= len(result)
        truncated |= len(result) < len(value)
        return result

    paragraphs = [dict(text=excerpt(p.text), style=excerpt(p.style.name))
                  for p in document.paragraphs[:100]]
    tables = []
    for table in document.tables[:20]:
        rows = []
        for row in table.rows[:20]:
            rows.append([excerpt(cell.text) for cell in row.cells[:20]])
            truncated |= len(row.cells) > 20
        tables.append(rows)
        truncated |= len(table.rows) > 20
    truncated |= len(document.paragraphs) > 100 or len(document.tables) > 20
    sections = []
    for section in document.sections[:100]:
        sections.append({key: getattr(section,key).inches if getattr(section,key) is not None else None
                         for key in ('page_width','page_height','left_margin','right_margin',
                                     'top_margin','bottom_margin')})
    truncated |= len(document.sections) > 100
    styles = []
    for style in document.styles:
        if not style.builtin:
            if len(styles) == 50:
                truncated = True
                break
            styles.append(dict(name=excerpt(style.name), font_name=excerpt(style.font.name or ''),
                               size_pt=style.font.size.pt if style.font.size else None,
                               bold=style.font.bold, italic=style.font.italic))
    return dict(paragraphs=paragraphs,tables=tables,sections=sections,styles=styles,
                inline_images=len(document.inline_shapes),
                page_breaks=len(document.element.xpath('.//w:br[@w:type="page"]')),
                truncated=truncated, visual_verification=False,
                limitations=['Structure only; no pagination or font-rendering verification.'])
