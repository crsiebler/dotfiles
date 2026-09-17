"""Original python-docx recipes for ordered document construction."""
import io
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt
from docx_support import DocumentError, fields, inspect_package, number, project_path, text

PAGE_FIELDS = {'width_inches':'page_width', 'height_inches':'page_height',
               'top_margin_inches':'top_margin', 'bottom_margin_inches':'bottom_margin',
               'left_margin_inches':'left_margin', 'right_margin_inches':'right_margin'}


def page_settings(section, data):
    fields(data, PAGE_FIELDS)
    for key, value in data.items():
        minimum = 0 if 'margin' in key else 0.1
        setattr(section, PAGE_FIELDS[key], Inches(number(value, minimum, 22)))
    if (section.left_margin + section.right_margin >= section.page_width
            or section.top_margin + section.bottom_margin >= section.page_height):
        raise DocumentError(2, 'input: margins leave no content area')


def add_styles(document, styles):
    if not isinstance(styles, dict) or len(styles) > 50:
        raise DocumentError(2, 'input: styles must be a map with at most 50 entries')
    for name, options in styles.items():
        text(name, 100)
        fields(options, ('font_name','size_pt','bold','italic','base'))
        if not name or name in document.styles:
            raise DocumentError(2, 'input: custom style name is empty or already exists')
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        base = document.styles[text(options.get('base', 'Normal'), 100)]
        if base.type != WD_STYLE_TYPE.PARAGRAPH:
            raise DocumentError(2, 'input: base must be a paragraph style')
        style.base_style = base
        if 'font_name' in options:
            style.font.name = text(options['font_name'], 100)
        if 'size_pt' in options:
            style.font.size = Pt(number(options['size_pt'], 1, 200))
        for key in ('bold','italic'):
            if key in options:
                if type(options[key]) is not bool:
                    raise DocumentError(2, 'input: bold/italic must be boolean')
                setattr(style.font, key, options[key])


def create_document(root, spec):
    fields(spec, ('blocks','page','styles'))
    blocks = spec.get('blocks')
    if not isinstance(blocks, list) or not 1 <= len(blocks) <= 1000:
        raise DocumentError(2, 'input: provide 1..1000 blocks')
    document = Document()
    section = document.sections[0]
    page_settings(section, dict(width_inches=8.5, height_inches=11,
                               top_margin_inches=1, bottom_margin_inches=1,
                               left_margin_inches=1, right_margin_inches=1))
    page_settings(section, spec.get('page', {}))
    add_styles(document, spec.get('styles', {}))
    expected_paragraphs, expected_tables = [], []
    image_count = break_count = char_count = cell_count = 0
    for block in blocks:
        if not isinstance(block, dict):
            raise DocumentError(2, 'input: block must be an object')
        kind = block.get('type')
        if kind in ('paragraph', 'heading'):
            fields(block, ('type','text','level') if kind == 'heading' else ('type','text','style'))
            content = text(block.get('text'))
            char_count += len(content)
            if kind == 'heading':
                level = block.get('level', 1)
                if type(level) is not int or not 0 <= level <= 9:
                    raise DocumentError(2, 'input: heading level must be 0..9')
                paragraph = document.add_heading(content, level)
            else:
                paragraph = document.add_paragraph(content, text(block.get('style','Normal'),100))
            expected_paragraphs.append((paragraph.text, paragraph.style.name))
        elif kind == 'table':
            fields(block, ('type','rows','style'))
            rows = block.get('rows')
            if (not isinstance(rows, list) or not 1 <= len(rows) <= 1000
                    or not isinstance(rows[0], list) or not 1 <= len(rows[0]) <= 100
                    or any(not isinstance(row,list) or len(row)!=len(rows[0]) for row in rows)):
                raise DocumentError(2, 'input: table must be a nonempty rectangle')
            cell_count += len(rows) * len(rows[0])
            if cell_count > 10000:
                raise DocumentError(2, 'input: total table cell limit exceeded')
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = text(block.get('style', 'Table Grid'), 100)
            for source, target in zip(rows, table.rows):
                for value, cell in zip(source, target.cells):
                    cell.text = text(value)
                    char_count += len(value)
            expected_tables.append(rows)
        elif kind == 'image':
            fields(block, ('type','path','width_inches'))
            image_count += 1
            if image_count > 100:
                raise DocumentError(2, 'input: image limit exceeded')
            width = Inches(number(block.get('width_inches'), 0.01, 22))
            if width > section.page_width - section.left_margin - section.right_margin:
                raise DocumentError(2, 'input: image exceeds section content width')
            path = project_path(root, block.get('path'))
            signature = path.read_bytes()[:8]
            if not (signature.startswith(b'\x89PNG\r\n\x1a\n') or signature.startswith(b'\xff\xd8\xff')):
                raise DocumentError(2, 'input: inline image must be PNG or JPEG')
            document.add_picture(str(path), width=width)
            expected_paragraphs.append(('', 'Normal'))
        elif kind == 'page_break':
            fields(block, ('type',))
            document.add_page_break()
            expected_paragraphs.append(('', 'Normal'))
            break_count += 1
        elif kind == 'section':
            fields(block, ('type','page'))
            section = document.add_section()
            page_settings(section, block.get('page', {}))
            expected_paragraphs.append(('', 'Normal'))
        else:
            raise DocumentError(2, 'input: unsupported block type')
        if char_count > 500000:
            raise DocumentError(2, 'input: total text limit exceeded')
    stream = io.BytesIO()
    document.save(stream)
    content = stream.getvalue()
    inspect_package(content)
    reopened = Document(io.BytesIO(content))
    if ([(p.text,p.style.name) for p in reopened.paragraphs] != expected_paragraphs
            or [[[c.text for c in row.cells] for row in t.rows] for t in reopened.tables] != expected_tables
            or len(reopened.inline_shapes) != image_count
            or len(reopened.sections) != len(document.sections)):
        raise DocumentError(5, 'validation: reopened document content differs')
    for original, actual in zip(document.sections, reopened.sections):
        if any(getattr(original, key) != getattr(actual, key) for key in PAGE_FIELDS.values()):
            raise DocumentError(5, 'validation: reopened page settings differ')
    for name in spec.get('styles', {}):
        before, after = document.styles[name], reopened.styles[name]
        if any(getattr(before.font, k) != getattr(after.font, k) for k in ('name','size','bold','italic')):
            raise DocumentError(5, 'validation: reopened style properties differ')
    if len(reopened.element.xpath('.//w:br[@w:type="page"]')) != break_count:
        raise DocumentError(5, 'validation: reopened page breaks differ')
    return content, dict(paragraphs=len(reopened.paragraphs), tables=len(reopened.tables),
                         inline_images=image_count, sections=len(reopened.sections), page_breaks=break_count)
