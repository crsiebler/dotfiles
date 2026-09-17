"""Bounded formula/cache evidence using independent read-only workbook views."""
from datetime import date, datetime, time, timedelta
import io
import openpyxl
from openpyxl.utils.cell import get_column_letter
from xlsx_input import ReadError
from xlsx_metadata import bounds, metadata


def extract(content, selected, region, max_cells, max_chars):
    formulas = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    caches = None
    try:
        caches = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        available = [s.title for s in formulas.worksheets]
        if any(len(name) > 31 for name in available):
            raise ReadError(2, 'input: worksheet name exceeds Excel limit')
        if selected and (len(set(selected)) > 20 or any(s not in available for s in selected)):
            raise ReadError(2, 'input: select at most 20 existing worksheets')
        names = [n for n in available if n in selected] if selected else available[:20]
        selected_bounds = bounds(region) if region else None
        info = metadata(content, names)
        remaining, characters = max_cells, max_chars
        truncated = not selected and len(available) > 20
        results = []

        def value(item):
            nonlocal characters, truncated
            if isinstance(item, (date, datetime, time)):
                item = item.isoformat()
            elif isinstance(item, timedelta):
                item = str(item)
            elif item is not None and not isinstance(item, (str, int, float, bool)):
                item = '[unsupported cell value]'
            if isinstance(item, str):
                clipped = item[:characters]
                characters -= len(clipped)
                truncated |= len(clipped) < len(item)
                return clipped
            return item

        for name in names:
            sheet, cache = formulas[name], caches[name]
            data = info[name]
            if selected_bounds:
                left, top, right, bottom = selected_bounds
            else:
                left, top = 1, 1
                right, bottom = min(data['max_column'], 100), min(data['max_row'], 2000)
                truncated |= right < data['max_column'] or bottom < data['max_row']
            requested = f'{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}'
            cells = []
            # Explicit bounds override unreliable source dimension hints.
            needed = (right-left+1)*(bottom-top+1)
            truncated |= needed > remaining
            if remaining:
                stop_row = min(bottom, top + (remaining-1)//(right-left+1))
                # Limit allocation for a single very wide requested row too.
                stop_col = min(right, left+remaining-1)
                limits = dict(min_row=top, max_row=stop_row, min_col=left, max_col=stop_col)
                for r, (row, cached_row) in enumerate(zip(sheet.iter_rows(**limits), cache.iter_rows(**limits)), top):
                    for c, (cell, cached_cell) in enumerate(zip(row, cached_row), left):
                        if not remaining:
                            break
                        remaining -= 1
                        entry = dict(coordinate=f'{get_column_letter(c)}{r}', type=cell.data_type,
                                     row_hidden=r in data['hidden_rows'],
                                     column_hidden=any(a <= c <= b for a, b in data['hidden_columns']),
                                     number_format=value(cell.number_format))
                        if cell.data_type == 'f':
                            entry.update(formula=value(cell.value), formula_supported=isinstance(cell.value, str),
                                         cached=value(cached_cell.value),
                                         cached_type=cached_cell.data_type,
                                         cache_status='missing' if cached_cell.value is None else 'present',
                                         cache_source='unknown')
                        else:
                            entry['value'] = value(cell.value)
                        cells.append(entry)
            results.append(dict(name=name, state=sheet.sheet_state, range=requested,
                                cells=cells, last_coordinate=cells[-1]['coordinate'] if cells else None))
        return dict(selected_sheets=names, sheets=results, truncated=bool(truncated), recalculated=False,
                    omitted_chart_sheets=len(formulas.chartsheets), limitations=[
                        'Present caches may be stale or placeholders; missing caches are not zero.',
                        'Unselected sheets/ranges are not inspected; empty extraction is not empty content.',
                        'Charts, images, comments, rich formatting, macros, external links and merged layout are omitted.',
                        'No save, calculation, external refresh or visual verification is performed.'])
    finally:
        formulas.close()
        if caches is not None:
            caches.close()
