"""Source-located bounded extraction without creating notes or modifying slides."""
import io
import re
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx_input import ReadError


def selection(value, count):
    if value is None:
        return list(range(1, min(count, 100) + 1)), count > 100
    selected = set()
    if len(value) > 1000:
        raise ReadError(2, 'input: slide selection too long')
    for part in value.split(','):
        if not re.fullmatch(r'[0-9]+(?:-[0-9]+)?', part):
            raise ReadError(2, 'input: expected slide numbers or inclusive ranges')
        ends = [int(v) for v in part.split('-')]
        first, last = ends[0], ends[-1]
        if not 1 <= first <= last <= count or last - first >= 100:
            raise ReadError(2, 'input: slide selection out of range or too large')
        selected.update(range(first, last + 1))
        if len(selected) > 100:
            raise ReadError(2, 'input: select at most 100 slides')
    return sorted(selected), False


class Extractor:
    def __init__(self, max_records, max_chars, max_values):
        self.records = []
        self.unsupported = []
        self.max_records = max_records
        self.characters = max_chars
        self.values = max_values
        self.truncated = False

    def text(self, value):
        value = str(value)
        result = value[:self.characters]
        self.characters -= len(result)
        self.truncated |= len(result) < len(value)
        return result

    def take(self, values):
        for value in values:
            if not self.values:
                self.truncated = True
                break
            self.values -= 1
            yield value

    def omit(self, slide, location, reason):
        if len(self.unsupported) < self.max_records:
            self.unsupported.append(dict(slide=slide, location=location, reason=reason))
        else:
            self.truncated = True

    def room(self):
        if len(self.records) >= self.max_records:
            self.truncated = True
            return False
        return True

    def shapes(self, shapes, slide, prefix, depth=0):
        for index, shape in enumerate(shapes, 1):
            if not self.room():
                break
            location = f'{prefix}/shape-{index}'
            base = dict(slide=slide, location=location)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                if depth >= 10:
                    self.omit(slide, location, 'group_depth')
                else:
                    self.shapes(shape.shapes, slide, location + '/group', depth + 1)
            elif shape.has_table:
                table = shape.table
                cells = (dict(row=r, column=c, text=cell.text)
                         for r, row in enumerate(table.rows, 1)
                         for c, cell in enumerate(row.cells, 1))
                extracted = [dict(cell, text=self.text(cell['text'])) for cell in self.take(cells)]
                self.records.append(dict(base, kind='table', cells=extracted,
                                         rows=len(table.rows), columns=len(table.columns)))
            elif shape.has_chart:
                self.chart(shape.chart, base)
            elif shape.has_text_frame:
                self.records.append(dict(base, kind='text', text=self.text(shape.text)))
            else:
                reason = 'picture' if shape.shape_type == MSO_SHAPE_TYPE.PICTURE else 'visual_or_unsupported_shape'
                self.omit(slide, location, reason)

    def chart(self, chart, base):
        try:
            if len(chart.plots) != 1:
                raise ValueError('multiple plots')
            categories = chart.plots[0].categories
            if categories.depth != 1:
                raise ValueError('nonflat categories')
            # Category-based series expose values; scatter/bubble do not provide
            # this supported combination of category labels and numeric values.
            series = list(chart.series)
            for item in series:
                item.values
        except (AttributeError, ValueError, TypeError):
            self.omit(base['slide'], base['location'], 'unsupported_chart')
            return
        labels = [self.text(category.label) for category in self.take(categories)]
        data = []
        for item in self.take(series):
            data.append(dict(name=self.text(item.name), values=list(self.take(item.values))))
        self.records.append(dict(base, kind='chart', categories=labels, series=data,
                                 values_source='stored_cache'))


def extract(content, selected, max_records, max_chars, max_values):
    presentation = Presentation(io.BytesIO(content))
    slides, truncated = selection(selected, len(presentation.slides))
    reader = Extractor(max_records, max_chars, max_values)
    reader.truncated = truncated
    for number in slides:
        slide = presentation.slides[number - 1]
        reader.shapes(slide.shapes, number, f'slide-{number}')
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is None:
                reader.omit(number, f'slide-{number}/notes', 'missing_notes_text_placeholder')
            elif reader.room():
                reader.records.append(dict(kind='notes', slide=number,
                                           location=f'slide-{number}/notes', text=reader.text(frame.text)))
    return dict(slide_count=len(presentation.slides), selected_slides=slides,
                records=reader.records, unsupported=reader.unsupported, truncated=reader.truncated,
                limitations=[
                    'Stored shape order is not visual reading order; unselected slides are not inspected.',
                    'Pictures and other visual-only content require visual inspection; no OCR.',
                    'Layout/master content, comments, animation, formatting and hyperlink targets are omitted.',
                    'Chart values are stored caches, not recalculated or refreshed.',
                    'Empty or truncated extraction does not establish empty visual content.'])
