"""Independent PPTX reader fixtures with source preservation and bounded output."""
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / 'ai/plugins/researching/skills/read-pptx/scripts/read_pptx.py'
PYTHON = os.environ.get('SKILL_TEST_PYTHON', sys.executable)


class ReadPptxTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=ROOT / 'tests')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        code = '''from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from PIL import Image
p=Presentation();layout=p.slide_layouts.get_by_name('Blank')
s=p.slides.add_slide(layout)
s.shapes.add_textbox(0,0,Inches(3),Inches(1)).text='Résumé slide one'
s.notes_slide.notes_text_frame.text='Speaker evidence'
t=s.shapes.add_table(2,2,0,Inches(1),Inches(4),Inches(1)).table
t.cell(0,0).text='Name';t.cell(0,1).text='Value';t.cell(1,0).text='A';t.cell(1,1).text='001'
s=p.slides.add_slide(layout)
g=s.shapes.add_group_shape();g.shapes.add_textbox(0,0,Inches(3),Inches(1)).text='Grouped evidence'
data=CategoryChartData();data.categories=['Q1','Q2'];data.add_series('Units',[12,18])
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,0,Inches(1),Inches(4),Inches(3),data)
Image.new('RGB',(2,2),'blue').save('image.png')
s.shapes.add_picture('image.png',Inches(5),0)
s=p.slides.add_slide(layout)
xy=XyChartData();series=xy.add_series('XY');series.add_data_point(1,2)
s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER,0,0,Inches(4),Inches(3),xy)
p.save('source.pptx')
'''
        result = subprocess.run([PYTHON, '-c', code], cwd=self.root, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)

    def cli(self, *args, script=SCRIPT, flags=()):
        return subprocess.run([PYTHON, *flags, str(script), '--project', str(self.root), *args],
                              cwd=self.root, capture_output=True, text=True)

    def test_notes_tables_chart_groups_locations_and_unchanged_source(self):
        source = self.root / 'source.pptx'; before = source.read_bytes()
        result = self.cli('read', 'source.pptx')
        self.assertEqual(result.returncode, 0, result.stderr)
        data = json.loads(result.stdout)
        self.assertEqual(data['selected_slides'], [1, 2, 3])
        self.assertEqual(data['sha256'], hashlib.sha256(before).hexdigest())
        self.assertEqual(source.read_bytes(), before)
        records = data['records']
        notes = next(r for r in records if r['kind'] == 'notes')
        self.assertEqual((notes['location'], notes['text']), ('slide-1/notes', 'Speaker evidence'))
        table = next(r for r in records if r['kind'] == 'table')
        self.assertEqual(table['cells'][-1], {'row': 2, 'column': 2, 'text': '001'})
        chart = next(r for r in records if r['kind'] == 'chart')
        self.assertEqual(chart['categories'], ['Q1', 'Q2'])
        self.assertEqual(chart['series'], [{'name': 'Units', 'values': [12.0, 18.0]}])
        grouped = next(r for r in records if r.get('text') == 'Grouped evidence')
        self.assertEqual(grouped['location'], 'slide-2/shape-1/group/shape-1')
        self.assertIn('picture', [u['reason'] for u in data['unsupported']])
        self.assertIn('unsupported_chart', [u['reason'] for u in data['unsupported']])
        self.assertFalse(data['truncated'])

    def test_selection_limits_and_failures(self):
        result = self.cli('read', 'source.pptx', '--slides', '2-3,2')
        self.assertEqual(result.returncode, 0, result.stderr)
        data = json.loads(result.stdout)
        self.assertEqual(data['selected_slides'], [2, 3])
        self.assertTrue(all(r['slide'] in (2, 3) for r in data['records']))
        self.assertNotIn('Speaker evidence', result.stdout)
        for value in ('0', '4', '3-1', 'a', '1,'):
            self.assertEqual(self.cli('read', 'source.pptx', '--slides', value).returncode, 2)
        result = self.cli('read', 'source.pptx', '--max-records', '2', '--max-chars', '5', '--max-values', '1')
        data = json.loads(result.stdout)
        self.assertTrue(data['truncated'])
        self.assertLessEqual(len(data['records']), 2)
        self.assertEqual(data['records'][0]['text'], 'Résum')
        self.assertEqual(self.cli('check', flags=('-S',)).returncode, 3)
        self.assertEqual(self.cli('check').returncode, 0)
        (self.root / 'bad.pptx').write_text('PRIVATE INPUT')
        result = self.cli('read', 'bad.pptx')
        self.assertEqual(result.returncode, 2)
        self.assertNotIn('PRIVATE', result.stderr)
        (self.root / 'link.pptx').symlink_to(self.root / 'source.pptx')
        self.assertEqual(self.cli('read', 'link.pptx').returncode, 2)
        self.assertEqual(self.cli('read', '../escape.pptx').returncode, 2)

    def test_isolated_reader_writes_no_source_or_resources(self):
        installed = self.root / 'installed' / 'read-pptx'
        shutil.copytree(SCRIPT.parents[1], installed, ignore=shutil.ignore_patterns('__pycache__'))
        before = {p.relative_to(self.root): p.read_bytes() for p in self.root.rglob('*') if p.is_file()}
        result = self.cli('read', 'source.pptx', script=installed / 'scripts/read_pptx.py')
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(before, {p.relative_to(self.root): p.read_bytes() for p in self.root.rglob('*') if p.is_file()})


if __name__ == '__main__':
    unittest.main()
