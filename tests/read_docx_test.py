"""Source-located DOCX reader regression fixtures, independent of create-docx."""
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

ROOT=Path(__file__).resolve().parents[1]
SCRIPT=ROOT/'ai/plugins/researching/skills/read-docx/scripts/read_docx.py'
PYTHON=os.environ.get('SKILL_TEST_PYTHON',sys.executable)


class ReadDocxTest(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir=ROOT/'tests')
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        code='''from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
d=Document()
d.add_heading('Résumé 世界',1)
t=d.add_table(rows=1,cols=2)
t.cell(0,0).text='A';t.cell(0,1).text='B'
nested=t.cell(0,0).add_table(rows=1,cols=1)
nested.cell(0,0).text='nested omitted'
d.add_paragraph('After table')
insert=OxmlElement('w:ins');p=OxmlElement('w:p');r=OxmlElement('w:r');text=OxmlElement('w:t')
text.text='UNACCEPTED REVISION';r.append(text);p.append(r);insert.append(p)
d.element.body.insert(3,insert)
s=d.sections[0]
s.header.paragraphs[0].text='Header evidence'
s.footer.paragraphs[0].text='Footer evidence'
s.first_page_header.paragraphs[0].text='Disabled first page header'
d.add_section()
d.add_paragraph('Last paragraph')
d.save('source.docx')
'''
        result=subprocess.run([PYTHON,'-c',code],cwd=self.root,capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)

    def cli(self,*args,script=SCRIPT,flags=()):
        return subprocess.run([PYTHON,*flags,str(script),'--project',str(self.root),*args],
                              cwd=self.root,capture_output=True,text=True)

    def test_order_locations_headers_revisions_and_preserved_source(self):
        source=self.root/'source.docx';before=hashlib.sha256(source.read_bytes()).hexdigest()
        result=self.cli('read','source.docx')
        self.assertEqual(result.returncode,0,result.stderr)
        data=json.loads(result.stdout)
        records=data['records']
        self.assertEqual([r['kind'] for r in records[:3]],['paragraph','table','paragraph'])
        self.assertEqual(records[0]['text'],'Résumé 世界')
        self.assertEqual(records[0]['heading_level'],1)
        self.assertEqual(records[1]['location'],'body/block-2')
        self.assertEqual(records[1]['cells'][0]['row'],1)
        self.assertEqual(records[1]['cells'][0]['column'],1)
        self.assertEqual(records[2]['text'],'After table')
        self.assertNotIn('UNACCEPTED REVISION',result.stdout)
        self.assertNotIn('nested omitted',result.stdout)
        self.assertEqual(data['unsupported']['tracked_insertions'],1)
        self.assertEqual(data['unsupported']['nested_tables'],1)
        second=next(c for c in data['contexts'] if c['section']==2 and c['variant']=='header')
        self.assertTrue(second['linked_to_previous'])
        self.assertEqual(second['definition_section'],1)
        first=next(c for c in data['contexts'] if c['section']==1 and c['variant']=='first_page_header')
        self.assertFalse(first['enabled'])
        self.assertIn('Header evidence',result.stdout)
        self.assertIn('Footer evidence',result.stdout)
        self.assertFalse(data['truncated'])
        self.assertEqual(before,hashlib.sha256(source.read_bytes()).hexdigest())
        self.assertEqual(data['sha256'],before)

    def test_omitted_grid_positions_and_merged_cells(self):
        code = """from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
d=Document()
t=d.add_table(rows=2,cols=3)
t.cell(0,0).text='omitted';t.cell(0,1).text='second';t.cell(0,2).text='third'
t.cell(1,0).merge(t.cell(1,1)).text='merged'
t.cell(1,2).text='last'
row=t.rows[0]._tr
row.remove(row.tc_lst[0])
before=OxmlElement('w:gridBefore');before.set(qn('w:val'),'1')
row.get_or_add_trPr().append(before)
d.save('grid.docx')
"""
        result=subprocess.run([PYTHON,'-c',code],cwd=self.root,capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)
        result=self.cli('read','grid.docx')
        self.assertEqual(result.returncode,0,result.stderr)
        cells=json.loads(result.stdout)['records'][0]['cells']
        self.assertEqual([(c['row'],c['column'],c['text']) for c in cells],
                         [(1,2,'second'),(1,3,'third'),(2,1,'merged'),(2,2,'merged'),(2,3,'last')])
        self.assertEqual(cells[0]['location'],'body/block-1/row-1/cell-2')

    def test_limits_and_invalid_paths_packages_dependencies(self):
        result=self.cli('read','source.docx','--max-records','2','--max-chars','5','--max-cells','1')
        self.assertEqual(result.returncode,0,result.stderr)
        data=json.loads(result.stdout)
        self.assertTrue(data['truncated'])
        self.assertLessEqual(len(data['records']),2)
        self.assertEqual(self.cli('read','source.docx','--max-records','0').returncode,2)
        (self.root/'link.docx').symlink_to(self.root/'source.docx')
        self.assertEqual(self.cli('read','link.docx').returncode,2)
        self.assertEqual(self.cli('read','../escape.docx').returncode,2)
        (self.root/'bad.docx').write_text('SECRET INVALID INPUT')
        result=self.cli('read','bad.docx')
        self.assertEqual(result.returncode,2)
        self.assertNotIn('SECRET',result.stderr)
        self.assertEqual(self.cli('check',flags=('-S',)).returncode,3)
        self.assertEqual(self.cli('check').returncode,0)

    def test_isolated_reader_needs_no_creation_skill_and_writes_nothing(self):
        installed=self.root/'installed'/'read-docx'
        shutil.copytree(SCRIPT.parents[1],installed,ignore=shutil.ignore_patterns('__pycache__'))
        before={p.relative_to(self.root):p.read_bytes() for p in self.root.rglob('*') if p.is_file()}
        result=self.cli('read','source.docx',script=installed/'scripts/read_docx.py')
        self.assertEqual(result.returncode,0,result.stderr)
        after={p.relative_to(self.root):p.read_bytes() for p in self.root.rglob('*') if p.is_file()}
        self.assertEqual(before,after)


if __name__=='__main__':
    unittest.main()
